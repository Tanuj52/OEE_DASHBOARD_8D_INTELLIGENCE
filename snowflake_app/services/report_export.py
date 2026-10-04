"""
8D report exports: PDF (reportlab), Word (python-docx) and an email-safe HTML
summary, all rendered from the same run snapshot.

Agent output is GitHub-flavoured Markdown (### headings, pipe tables, lists,
**bold**). parse_blocks() turns it into a small block model so each renderer
lays it out natively instead of shipping raw Markdown.
"""

import io
import os
import re
from datetime import datetime
from html import escape

from snowflake_app.core.monitoring import fmt_hours

STATIC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static")

ORANGE, INK, MUTED, SUBTLE, LINE, SOFT = "#F26B1D", "#1B1816", "#57524D", "#78716B", "#E5E2DE", "#F5F4F2"

# Highlights pulled from specific agent sections: (key, step id, heading prefix, label)
HIGHLIGHTS = [
    ("summary", "d8", "executive summary", "Executive summary"),
    ("root_cause", "d4", "verified root cause", "Verified root cause"),
    ("containment", "d3", "containment actions", "Containment"),
    ("copq", "d8", "cost of poor quality", "Cost of poor quality"),
    ("pdm", "predictive", "decision", "Maintenance decision"),
]


# ---------------------------------------------------------------------------
# Markdown -> blocks
# ---------------------------------------------------------------------------

_HEAD = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_ITEM = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|(?<![\w*])\*(?!\s)[^*]+?\*(?![\w*]))")


def clean_heading(text: str) -> str:
    """Drop prompt scaffolding the model sometimes echoes, e.g. "6M analysis (table: factor | evidence)"."""
    return re.sub(r"\s*\((?:table|list):[^)]*\)\s*$", "", text, flags=re.I).strip()


def clean_md(text: str | None) -> str:
    return re.sub(r"^(#{1,6}\s+.*)$", lambda m: clean_heading(m.group(1)), text or "", flags=re.M)


def _cells(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_blocks(text: str | None) -> list[dict]:
    """Blocks: heading{level,text} · para{text} · list{ordered,items:[(depth,text)]} · table{header,rows} · quote{text}."""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    out, para, i = [], [], 0

    def flush():
        if para:
            out.append({"t": "para", "text": " ".join(p.strip() for p in para)})
            para.clear()

    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            flush()
            i += 1
            continue
        m = _HEAD.match(ln)
        if m:
            flush()
            out.append({"t": "heading", "level": len(m.group(1)), "text": clean_heading(m.group(2))})
            i += 1
            continue
        if ln.lstrip().startswith("|") and i + 1 < len(lines) and _SEP.match(lines[i + 1]):
            flush()
            header = [h[:1].upper() + h[1:] for h in _cells(ln)]
            rows, i = [], i + 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                r = _cells(lines[i])
                rows.append((r + [""] * len(header))[:len(header)])
                i += 1
            out.append({"t": "table", "header": header, "rows": rows})
            continue
        m = _ITEM.match(ln)
        if m:
            flush()
            ordered = m.group(2)[0].isdigit()
            items = []
            while i < len(lines):
                m = _ITEM.match(lines[i])
                if m:
                    items.append([min(len(m.group(1).expandtabs(4)) // 2, 3), m.group(3).strip()])
                elif lines[i].strip() and lines[i].startswith(" ") and items:
                    items[-1][1] += " " + lines[i].strip()  # wrapped continuation line
                else:
                    break
                i += 1
            out.append({"t": "list", "ordered": ordered, "items": [tuple(x) for x in items]})
            continue
        if ln.lstrip().startswith(">"):
            flush()
            q = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                q.append(lines[i].lstrip()[1:].strip())
                i += 1
            out.append({"t": "quote", "text": " ".join(q)})
            continue
        if re.match(r"^\s*([-*_])\s*\1\s*\1[\s\1]*$", ln):
            flush()
            i += 1
            continue
        para.append(ln)
        i += 1
    flush()
    return out


def inline_runs(text: str) -> list[tuple[str, str]]:
    """[(style, text)] with style in {"", "b", "i", "code"}."""
    runs = []
    for part in _INLINE.split(text or ""):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**") and len(part) > 4) or (part.startswith("__") and part.endswith("__") and len(part) > 4):
            runs.append(("b", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append(("code", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append(("i", part[1:-1]))
        else:
            runs.append(("", part))
    return runs


def section_markdown(output: str | None, heading_prefix: str) -> str:
    """Body of the first ### section whose heading starts with heading_prefix (case-insensitive)."""
    lines, grab, out = (output or "").split("\n"), False, []
    for ln in lines:
        m = _HEAD.match(ln)
        if m:
            if grab:
                break
            grab = clean_heading(m.group(2)).lower().startswith(heading_prefix)
            continue
        if grab:
            out.append(ln)
    return "\n".join(out).strip()


# ---------------------------------------------------------------------------
# Report model shared by every format
# ---------------------------------------------------------------------------

def _fmt_dt(iso: str | None) -> str:
    if not iso:
        return "—"
    try:
        return datetime.fromisoformat(iso).strftime("%b %-d, %Y · %I:%M %p")
    except ValueError:
        return iso


def _duration(snap: dict) -> str:
    try:
        s = int((datetime.fromisoformat(snap["finished_at"]) - datetime.fromisoformat(snap["created_at"])).total_seconds())
    except (KeyError, TypeError, ValueError):
        return "—"
    return f"{s // 60}m {s % 60:02d}s" if s >= 60 else f"{s}s"


def _signal(t: dict) -> str:
    if not t.get("sensor"):
        return "—"
    from snowflake_app.ui import kit as K  # display formatting only; no Streamlit calls

    return f"{K.sensor_label(t['sensor'])} {K.sensor_val(t['sensor'], t.get('value'))} {K.sensor_unit(t['sensor'])}".strip()


def build(snap: dict) -> dict:
    t = snap["trigger"]
    steps = {s["id"]: s for s in snap["steps"]}
    done = [s for s in snap["steps"] if s["status"] == "done"]
    highlights = []
    for key, sid, prefix, label in HIGHLIGHTS:
        body = section_markdown((steps.get(sid) or {}).get("output"), prefix)
        if body:
            highlights.append({"key": key, "label": label, "markdown": body, "step": (steps.get(sid) or {}).get("code")})
    rul = t.get("rul_hours")
    facts = [
        ("Machine", f"{t.get('machine_id') or '—'} · {t.get('machine_name') or ''}".strip(" ·")),
        ("Line", t.get("line") or "—"),
        ("Trigger", {"live": "Live monitoring alert", "fleet": "Fleet risk scan", "machine": "Machine selection"}.get(
            t.get("source"), (t.get("source") or "—").capitalize())),
        ("Signal", _signal(t)),
        ("Risk", (t.get("risk") or "—").capitalize()),
        ("Remaining useful life", "—" if rul is None else ("Limit crossed" if rul <= 0 else fmt_hours(rul))),
        ("Started", _fmt_dt(snap.get("created_at"))),
        ("Duration", _duration(snap)),
        ("Status", (snap.get("status") or "—").capitalize()),
        ("Outcome", (snap.get("decision") or {}).get("message") or "—"),
        ("AI model", f"{snap['provider'].get('label')} · {snap['provider'].get('model')}"),
        ("Agents completed", f"{snap['progress']['done']} of {snap['progress']['total']}"),
    ]
    return {
        "id": snap["id"],
        "title": f"{t.get('machine_id') or 'Plant'} {t.get('machine_name') or ''}".strip(),
        "subtitle": " · ".join(x for x in (t.get("line"), _fmt_dt(snap.get("created_at"))) if x),
        "facts": facts,
        "highlights": highlights,
        "sections": [{
            "id": s["id"], "code": s["code"], "title": s["title"], "subtitle": s["subtitle"],
            "evidence": [(tl["label"], tl["summary"]) for tl in s.get("tools", []) if tl.get("summary")],
            "blocks": parse_blocks(s.get("output")),
        } for s in done],
        "skipped": [s["code"] for s in snap["steps"] if s["status"] == "skipped"],
        "tokens": snap.get("total_tokens"),
        "actions": _actions(snap["id"]),
    }


def _actions(run_id: str) -> dict | None:
    """The report's tracked actions as a table block (services/capa_ops), so every export shows their status."""
    try:
        from snowflake_app.core import capa as C
        from snowflake_app.services import capa_ops
        acts = capa_ops.actions(run_id)
    except Exception:
        return None
    if not acts:
        return None
    rows = [[a["discipline"], a["action"], a.get("owner_name") or "—", _fmt_dt(a.get("due_at")).split(" · ")[0],
             C.STATUS_LABEL.get(a["status"], a["status"])] for a in acts]
    return {"t": "table", "header": ["Step", "Action", "Owner", "Due", "Status"], "rows": rows,
            "status": C.run_status(acts, datetime.now().timestamp())}


def file_stem(snap: dict) -> str:
    mid = re.sub(r"[^A-Za-z0-9-]", "", str(snap["trigger"].get("machine_id") or ""))
    return f"{snap['id']}{'_' + mid if mid else ''}"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONTS = False


def _register_fonts():
    global _FONTS
    if _FONTS:
        return
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, f in (("Inter", "Inter-Regular.ttf"), ("Inter-Medium", "Inter-Medium.ttf"), ("Inter-SemiBold", "Inter-SemiBold.ttf"),
                    ("Inter-Bold", "Inter-Bold.ttf"), ("Mono", "JetBrainsMono-Regular.ttf")):
        pdfmetrics.registerFont(TTFont(name, os.path.join(STATIC, f)))
    # Inter ships without italics here: map italic to regular so <i> never fails.
    addMapping("Inter", 0, 0, "Inter")
    addMapping("Inter", 1, 0, "Inter-SemiBold")
    addMapping("Inter", 0, 1, "Inter")
    addMapping("Inter", 1, 1, "Inter-SemiBold")
    _FONTS = True


def _rl_inline(text: str) -> str:
    parts = []
    for style, s in inline_runs(text):
        s = escape(s, quote=False)
        if style == "b":
            parts.append(f"<b>{s}</b>")
        elif style == "i":
            parts.append(f'<font color="{MUTED}">{s}</font>')
        elif style == "code":
            parts.append(f'<font name="Mono" size="8.5">{s}</font>')
        else:
            parts.append(s)
    return "".join(parts)


def to_pdf(snap: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (CondPageBreak, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    _register_fonts()
    rep = build(snap)
    c = colors.HexColor
    W, H = A4
    margin = 18 * mm
    width = W - 2 * margin

    base = ParagraphStyle("base", fontName="Inter", fontSize=9.5, leading=14, textColor=c(INK), alignment=TA_LEFT)
    st = {
        "body": base,
        "small": ParagraphStyle("small", parent=base, fontSize=8, leading=11, textColor=c(SUBTLE)),
        "label": ParagraphStyle("label", parent=base, fontName="Inter-SemiBold", fontSize=7, leading=9, textColor=c(SUBTLE)),
        "value": ParagraphStyle("value", parent=base, fontName="Inter-Medium", fontSize=9, leading=12),
        "h1": ParagraphStyle("h1", parent=base, fontName="Inter-Bold", fontSize=20, leading=25),
        "sec": ParagraphStyle("sec", parent=base, fontName="Inter-Bold", fontSize=13.5, leading=18),
        "secsub": ParagraphStyle("secsub", parent=base, fontSize=9, leading=12, textColor=c(SUBTLE)),
        "h3": ParagraphStyle("h3", parent=base, fontName="Inter-SemiBold", fontSize=10.5, leading=14, spaceBefore=8, spaceAfter=3),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.3, leading=11),
        "head": ParagraphStyle("head", parent=base, fontName="Inter-SemiBold", fontSize=8, leading=10.5, textColor=colors.white),
        "ev": ParagraphStyle("ev", parent=base, fontSize=8.3, leading=11.5, textColor=c(MUTED)),
        "hl_label": ParagraphStyle("hl_label", parent=base, fontName="Inter-SemiBold", fontSize=8, leading=10, textColor=c(ORANGE)),
        "toc": ParagraphStyle("toc", parent=base, fontSize=9, leading=15),
    }

    def blocks(bl: list[dict], compact: bool = False) -> list:
        out = []
        for b in bl:
            if b["t"] == "heading":
                out.append(Paragraph(_rl_inline(b["text"]), st["h3"]))
            elif b["t"] == "para":
                out += [Paragraph(_rl_inline(b["text"]), st["body"]), Spacer(1, 4)]
            elif b["t"] == "quote":
                out += [Paragraph(_rl_inline(b["text"]), st["ev"]), Spacer(1, 4)]
            elif b["t"] == "list":
                items = [ListItem(Paragraph(_rl_inline(tx), st["body"]), leftIndent=12 + 12 * d, bulletColor=c(ORANGE))
                         for d, tx in b["items"]]
                out += [ListFlowable(items, bulletType="1" if b["ordered"] else "bullet", start="1" if b["ordered"] else None,
                                     bulletFontName="Inter-SemiBold", bulletFontSize=8.5, leftIndent=12, bulletColor=c(ORANGE)),
                        Spacer(1, 4)]
            elif b["t"] == "table":
                n = len(b["header"])
                lens = [max([len(b["header"][j])] + [len(r[j]) for r in b["rows"]]) for j in range(n)]
                tot = sum(min(x, 60) + 6 for x in lens)
                widths = [width * (min(x, 60) + 6) / tot for x in lens]
                data = [[Paragraph(_rl_inline(h), st["head"]) for h in b["header"]]] + \
                       [[Paragraph(_rl_inline(x), st["cell"]) for x in r] for r in b["rows"]]
                tbl = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), c(ORANGE)),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, c("#FAF9F7")]),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.4, c(LINE)),
                    ("BOX", (0, 0), (-1, -1), 0.5, c(LINE)),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]))
                out += [tbl, Spacer(1, 6)]
        return out

    story = []
    # ---- title block
    story += [Paragraph("8D PROBLEM-SOLVING REPORT", ParagraphStyle("kick", parent=st["label"], textColor=c(ORANGE), fontSize=8)),
              Spacer(1, 4), Paragraph(escape(rep["title"]), st["h1"]),
              Paragraph(f'<font name="Mono">{escape(rep["id"])}</font> · {escape(rep["subtitle"])}', st["secsub"]), Spacer(1, 10)]
    cols = 3
    cells = [[Paragraph(k.upper(), st["label"]), Paragraph(escape(str(v)), st["value"])] for k, v in rep["facts"]]
    grid = []
    for r in range(0, len(cells), cols):
        row = cells[r:r + cols] + [[""]] * (cols - len(cells[r:r + cols]))
        grid.append(row)
    ft = Table(grid, colWidths=[width / cols] * cols, hAlign="LEFT")
    ft.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c(SOFT)), ("BOX", (0, 0), (-1, -1), 0.5, c(LINE)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story += [ft, Spacer(1, 14)]

    # ---- highlights
    for h in rep["highlights"][:4]:
        inner = [Paragraph(f'{escape(h["label"].upper())} <font color="{SUBTLE}">· {escape(h["step"] or "")}</font>', st["hl_label"]),
                 Spacer(1, 3)] + blocks(parse_blocks(h["markdown"]))
        box = Table([[inner]], colWidths=[width], hAlign="LEFT")
        box.setStyle(TableStyle([("LINEBEFORE", (0, 0), (0, -1), 2.2, c(ORANGE)), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                                 ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story += [box, Spacer(1, 10)]

    # ---- action status (tracked in the app since the report finished)
    if rep["actions"]:
        stt = rep["actions"]["status"]
        story += [Paragraph(f'ACTION STATUS <font color="{SUBTLE}">· {escape(stt["label"])} · {stt.get("closed", 0)} of '
                            f'{stt["total"]} closed</font>', st["hl_label"]), Spacer(1, 4)] + blocks([rep["actions"]])

    # ---- contents
    story += [CondPageBreak(60 * mm), Paragraph("Contents", st["h3"])]
    toc = [[Paragraph(f'<font name="Inter-SemiBold" color="{ORANGE}">{escape(s["code"])}</font>', st["toc"]),
            Paragraph(f'{escape(s["title"])} <font color="{SUBTLE}">— {escape(s["subtitle"])}</font>', st["toc"])]
           for s in rep["sections"]]
    if rep["skipped"]:
        toc.append([Paragraph("", st["toc"]), Paragraph(f'<font color="{SUBTLE}">Not run: {", ".join(rep["skipped"])}</font>', st["toc"])])
    tt = Table(toc, colWidths=[16 * mm, width - 16 * mm], hAlign="LEFT")
    tt.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.3, c(LINE)), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [tt]

    # ---- sections
    for s in rep["sections"]:
        code_style = ParagraphStyle("code", parent=st["sec"], textColor=colors.white, alignment=1,
                                    fontSize=13.5 if len(s["code"]) <= 3 else 11)
        head = Table([[Paragraph(escape(s["code"]), code_style),
                       [Paragraph(escape(s["title"]), st["sec"]), Paragraph(escape(s["subtitle"]), st["secsub"])]]],
                     colWidths=[18 * mm, width - 18 * mm], hAlign="LEFT")
        head.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), c(ORANGE)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                  ("LEFTPADDING", (1, 0), (1, 0), 10), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LINEBELOW", (0, 0), (-1, 0), 0.6, c(ORANGE))]))
        ev = []
        if s["evidence"]:
            et = Table([[Paragraph(f'<font name="Inter-SemiBold" color="{INK}">Evidence · {escape(lb)}:</font> {escape(sm)}', st["ev"])]
                        for lb, sm in s["evidence"]], colWidths=[width], hAlign="LEFT")
            et.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c(SOFT)), ("LINEBEFORE", (0, 0), (0, -1), 2, c(ORANGE)),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            ev = [et, Spacer(1, 6)]
        body = blocks(s["blocks"])
        story += [CondPageBreak(45 * mm), Spacer(1, 12), KeepTogether([head, Spacer(1, 8)] + ev + body[:2])] + body[2:]

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(c(ORANGE))
        canvas.rect(0, H - 5, W, 5, stroke=0, fill=1)
        canvas.setFont("Inter", 7.5)
        canvas.setFillColor(c(SUBTLE))
        if doc.page > 1:
            canvas.drawString(margin, H - 12 * mm, f"8D Report {rep['id']}")
            canvas.drawRightString(W - margin, H - 12 * mm, rep["title"])
        canvas.setStrokeColor(c(LINE))
        canvas.line(margin, 13 * mm, W - margin, 13 * mm)
        canvas.drawString(margin, 9 * mm, f"8D Intelligence · generated with {snap['provider'].get('label')} "
                                          f"({snap['provider'].get('model')}) · {datetime.now():%b %-d, %Y}")
        canvas.drawRightString(W - margin, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=f"8D Report {rep['id']} — {rep['title']}", author="8D Intelligence",
                            subject="8D problem-solving report", creator="8D Intelligence")
    doc.build(story, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def to_docx(snap: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    rep = build(snap)
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.color.rgb = "Calibri", Pt(size), rgb(INK)
    doc.styles["Heading 1"].font.color.rgb = rgb(ORANGE)

    def shade(cell, fill):
        tc = cell._tc.get_or_add_tcPr()
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:color"), "auto")
        el.set(qn("w:fill"), fill.lstrip("#"))
        tc.append(el)

    def runs(par, text, size=None, color=None, bold=False):
        for style, s in inline_runs(text):
            r = par.add_run(s)
            r.bold = bold or style == "b"
            r.italic = style == "i"
            if style == "code":
                r.font.name = "Consolas"
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
        return par

    def blocks(bl):
        for b in bl:
            if b["t"] == "heading":
                doc.add_heading(b["text"], level=3)
            elif b["t"] == "para":
                runs(doc.add_paragraph(), b["text"])
            elif b["t"] == "quote":
                runs(doc.add_paragraph(style="Quote"), b["text"])
            elif b["t"] == "list":
                for d, tx in b["items"]:
                    base = "List Number" if b["ordered"] else "List Bullet"
                    runs(doc.add_paragraph(style=base if d == 0 else f"{base} {min(d + 1, 3)}"), tx)
            elif b["t"] == "table":
                t = doc.add_table(rows=1, cols=len(b["header"]))
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.LEFT
                for j, h in enumerate(b["header"]):
                    cell = t.rows[0].cells[j]
                    cell.text = ""
                    runs(cell.paragraphs[0], h, size=9, color="#FFFFFF", bold=True)
                    shade(cell, ORANGE)
                for r in b["rows"]:
                    cells = t.add_row().cells
                    for j, x in enumerate(r):
                        cells[j].text = ""
                        runs(cells[j].paragraphs[0], x, size=9)
                doc.add_paragraph()

    kick = doc.add_paragraph()
    runs(kick, "8D PROBLEM-SOLVING REPORT", size=8.5, color=ORANGE, bold=True)
    doc.add_heading(rep["title"], level=0)
    runs(doc.add_paragraph(), f"{rep['id']} · {rep['subtitle']}", size=9.5, color=SUBTLE)

    ft = doc.add_table(rows=0, cols=2)
    ft.style = "Table Grid"
    for k, v in rep["facts"]:
        cells = ft.add_row().cells
        cells[0].text, cells[1].text = "", ""
        runs(cells[0].paragraphs[0], k, size=9, color=MUTED, bold=True)
        runs(cells[1].paragraphs[0], str(v), size=9.5)
        shade(cells[0], SOFT)
    doc.add_paragraph()

    if rep["highlights"]:
        doc.add_heading("Summary", level=1)
        for h in rep["highlights"][:4]:
            doc.add_heading(f"{h['label']} ({h['step']})", level=2)
            blocks(parse_blocks(h["markdown"]))

    if rep["actions"]:
        stt = rep["actions"]["status"]
        doc.add_heading("Action status", level=1)
        runs(doc.add_paragraph(), f"{stt['label']} · {stt.get('closed', 0)} of {stt['total']} actions closed", size=9.5, color=SUBTLE)
        blocks([rep["actions"]])

    for s in rep["sections"]:
        doc.add_page_break()
        doc.add_heading(f"{s['code']} — {s['title']}", level=1)
        runs(doc.add_paragraph(), s["subtitle"], size=9.5, color=SUBTLE)
        for lb, sm in s["evidence"]:
            p = doc.add_paragraph(style="Quote")
            runs(p, f"**Evidence · {lb}:** {sm}", size=9)
        blocks(s["blocks"])

    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    runs(fp, f"8D Report {rep['id']} · {rep['title']} · page ", size=8, color=SUBTLE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    doc.core_properties.title = f"8D Report {rep['id']} — {rep['title']}"
    doc.core_properties.author = "8D Intelligence"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Email HTML (table layout, inline styles: renders in Outlook and Gmail)
# ---------------------------------------------------------------------------

def _h_inline(text: str) -> str:
    out = []
    for style, s in inline_runs(text):
        s = escape(s, quote=False)
        out.append({"b": f"<b>{s}</b>", "i": f"<i>{s}</i>",
                    "code": f'<code style="font-family:Consolas,monospace;font-size:12px">{s}</code>'}.get(style, s))
    return "".join(out)


def blocks_html(bl: list[dict]) -> str:
    p = f"margin:0 0 10px;font-size:14px;line-height:21px;color:{INK}"
    out = []
    for b in bl:
        if b["t"] == "heading":
            out.append(f'<p style="margin:14px 0 6px;font-size:14px;font-weight:700;color:{INK}">{_h_inline(b["text"])}</p>')
        elif b["t"] in ("para", "quote"):
            out.append(f'<p style="{p}">{_h_inline(b["text"])}</p>')
        elif b["t"] == "list":
            tag = "ol" if b["ordered"] else "ul"
            out.append(f'<{tag} style="margin:0 0 10px;padding-left:20px">' + "".join(
                f'<li style="{p};margin:0 0 4px;margin-left:{d * 14}px">{_h_inline(tx)}</li>' for d, tx in b["items"]) + f"</{tag}>")
        elif b["t"] == "table":
            th = "".join(f'<th align="left" style="background:{ORANGE};color:#FFFFFF;font-size:12px;padding:6px 8px">{_h_inline(h)}</th>'
                         for h in b["header"])
            trs = "".join("<tr>" + "".join(f'<td style="font-size:12.5px;padding:6px 8px;border-bottom:1px solid {LINE};'
                                           f'vertical-align:top;color:{INK}">{_h_inline(x)}</td>' for x in r) + "</tr>" for r in b["rows"])
            out.append(f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-collapse:collapse;'
                       f'border:1px solid {LINE};margin:0 0 12px"><tr>{th}</tr>{trs}</table>')
    return "".join(out)


def email_html(snap: dict, link: str | None, expires: str | None, note: str = "", sender: str = "",
               include_full: bool = False, attached: bool = False) -> str:
    rep = build(snap)
    facts = [f for f in rep["facts"] if f[0] in ("Machine", "Line", "Trigger", "Signal", "Risk", "Outcome", "Started", "Status")]
    fact_rows = "".join(
        f'<tr><td style="padding:5px 0;font-size:12px;color:{SUBTLE};width:150px;vertical-align:top">{escape(k)}</td>'
        f'<td style="padding:5px 0;font-size:13px;color:{INK};font-weight:600">{escape(str(v))}</td></tr>' for k, v in facts)
    hl = "".join(
        f'<tr><td style="padding:0 0 16px"><div style="border-left:3px solid {ORANGE};padding:2px 0 2px 14px">'
        f'<p style="margin:0 0 6px;font-size:11px;font-weight:700;letter-spacing:.06em;color:{ORANGE}">{escape(h["label"].upper())}</p>'
        f'{blocks_html(parse_blocks(h["markdown"]))}</div></td></tr>' for h in rep["highlights"][:3])
    if rep["actions"]:
        hl += (f'<tr><td style="padding:0 0 16px"><p style="margin:0 0 6px;font-size:11px;font-weight:700;letter-spacing:.06em;'
               f'color:{ORANGE}">ACTIONS · {escape(rep["actions"]["status"]["label"].upper())}</p>{blocks_html([rep["actions"]])}</td></tr>')
    note_html = (f'<tr><td style="padding:0 0 18px"><div style="background:{SOFT};border-radius:8px;padding:12px 14px;font-size:14px;'
                 f'line-height:21px;color:{INK}">{escape(note).replace(chr(10), "<br>")}'
                 + (f'<div style="margin-top:6px;font-size:12px;color:{SUBTLE}">— {escape(sender)}</div>' if sender else "")
                 + "</div></td></tr>") if note.strip() else ""
    cta = (f'<tr><td style="padding:4px 0 6px"><a href="{escape(link)}" style="display:inline-block;background:{ORANGE};color:#FFFFFF;'
           f'text-decoration:none;font-weight:700;font-size:14px;padding:12px 22px;border-radius:8px">Download the PDF report</a></td></tr>'
           f'<tr><td style="padding:0 0 22px;font-size:12px;color:{SUBTLE}">Secure link, valid until {escape(expires or "")}.</td></tr>'
           ) if link else (f'<tr><td style="padding:0 0 20px"><div style="border:1px solid {LINE};border-radius:8px;padding:10px 14px;'
                           f'font-size:13px;color:{INK}"><b>The full report is attached</b> as a PDF.</div></td></tr>' if attached else "")
    full = ""
    if include_full:
        full = "".join(
            f'<tr><td style="padding:18px 0 6px;border-top:1px solid {LINE}"><p style="margin:0;font-size:16px;font-weight:700;color:{INK}">'
            f'<span style="color:{ORANGE}">{escape(s["code"])}</span> &nbsp;{escape(s["title"])}</p>'
            f'<p style="margin:2px 0 10px;font-size:12px;color:{SUBTLE}">{escape(s["subtitle"])}</p>{blocks_html(s["blocks"])}</td></tr>'
            for s in rep["sections"])
    return f"""<!doctype html><html><body style="margin:0;padding:0;background:#F3F2F0">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#F3F2F0;padding:24px 12px"><tr><td align="center">
<table role="presentation" width="640" cellpadding="0" cellspacing="0" style="max-width:640px;width:100%;background:#FFFFFF;border-radius:12px;
 overflow:hidden;font-family:Segoe UI,Helvetica,Arial,sans-serif">
<tr><td style="background:{ORANGE};padding:20px 28px">
 <p style="margin:0;font-size:11px;font-weight:700;letter-spacing:.08em;color:#FFFFFF;opacity:.9">8D PROBLEM-SOLVING REPORT</p>
 <p style="margin:6px 0 0;font-size:22px;font-weight:700;color:#FFFFFF">{escape(rep["title"])}</p>
 <p style="margin:4px 0 0;font-size:13px;color:#FFFFFF">{escape(rep["id"])} · {escape(rep["subtitle"])}</p></td></tr>
<tr><td style="padding:24px 28px 8px"><table role="presentation" width="100%" cellpadding="0" cellspacing="0">
{note_html}
<tr><td style="padding:0 0 18px"><table role="presentation" width="100%" cellpadding="0" cellspacing="0">{fact_rows}</table></td></tr>
{cta}{hl}{full}
</table></td></tr>
<tr><td style="padding:16px 28px;background:{SOFT};font-size:11.5px;line-height:17px;color:{SUBTLE}">
 Sent from 8D Intelligence on Snowflake. Generated by {escape(snap["provider"].get("label") or "")} ({escape(snap["provider"].get("model") or "")}).
 {len(rep["sections"])} agent sections{f", {rep['tokens']:,} tokens" if rep.get("tokens") else ""}.</td></tr>
</table></td></tr></table></body></html>"""
